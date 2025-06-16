from openpyxl import load_workbook
import json

from util.ExcelUtil import ExcelUtil
from util.ObjUtil import ObjUtil

"""
用于解析词条页，将数据保存至resource/entry/词条.json
src/resource/tmp/entry_dict.json用于存储json对象和单元格位置的关系
"""

# 加载 Excel 文件
wb = load_workbook('../mnt/支援卡-6.9新卡追加.xlsx')

# 获取词条工作表
sheet = wb['词条']


def func():
    # 初始化列表来存储数据
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(list(row))
    print(data)
    # 重新设置列名
    data[0] = ['事件', 'SSR_事件', 'SSR_30', 'SSR_35', 'SSR_40', 'SSR_空', 'SSR_空', 'SR_事件', 'SR_22', 'SR_26',
               'SR_30', 'SR_空', 'SR_空']

    # 从第 2 行(index=1)开始加载数据
    data = data[1:]

    # 处理合并单元格
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left_value = sheet.cell(row=min_row, column=min_col).value
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if row - 1 < len(data) and col - 1 < len(data[row - 1]):
                    data[row - 1][col - 1] = top_left_value
    result = {
        "SSR": {},
        "SR": {}
    }

    # 提取 SSR 数据
    ssr_start_index = 1
    for row in data:
        event_type = row[0]
        # 将 None 值（空格）替换为 -1
        ssr_values = [value if value is not None else -1 for value in row[ssr_start_index + 1:ssr_start_index + 6]]
        result["SSR"][event_type] = ssr_values

    # 提取 SR 数据
    sr_start_index = 7
    for row in data:
        event_type = row[0]
        # 将 None 值（空格）替换为 -1
        sr_values = [value if value is not None else -1 for value in row[sr_start_index + 1:sr_start_index + 6]]
        result["SR"][event_type] = sr_values

    # 假设result是之前处理得到的包含SSR和SR数据的字典
    keys_to_process = []
    for rarity in ["SSR", "SR"]:
        for event_type in result[rarity].keys():
            if event_type == "删卡":
                keys_to_process.append((rarity, event_type))

    for rarity, event_type in keys_to_process:
        values = result[rarity][event_type]
        # 初始化新的字段
        result[rarity]["删卡"] = []
        result[rarity]["次限删卡"] = []

        # 初始化标志变量
        is_next_limit = False

        for value in values:
            if value == "次限":
                is_next_limit = True
            elif is_next_limit:
                result[rarity]["次限删卡"].append(value)
            else:
                result[rarity]["删卡"].append(value)

    with open('../resource/entry/词条.json', 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=4)


def func_dict():
    route_dict = {}
    for row in range(2, sheet.max_row + 1):
        cell_value = sheet.cell(row=row, column=1).value
        if ObjUtil.is_empty(cell_value):
            break
        if cell_value == '删卡':
            for col in range(3, 3 + 2):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SSR']['{cell_value}'][{col - 3}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
            for col in range(3 + 3, 3 + 5):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SSR']['次限{cell_value}'][{col - 3 - 3}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
        elif cell_value == '百分比' or cell_value == '基础值':
            for col in range(3, 3 + 5):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SSR']['{cell_value}'][{col - 3}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
            for col_left in range(3, 3 + 5):
                for col_right in range(col_left + 1, 3 + 5):
                    name1 = ExcelUtil.col_row_to_excel_col(col_left, row)
                    name2 = ExcelUtil.col_row_to_excel_col(col_right, row)
                    value = f"entry['SSR']['{cell_value}'][{col_left - 3}:{col_right - 3+1}]"
                    route_dict.update({f"词条!{name1}:{name2}": f"{value}"})
                    route_dict.update({f"'词条'!{name1}:{name2}": f"{value}"})
        else:
            for col in range(3, 3 + 5):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SSR']['{cell_value}'][{col - 3}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
    for row in range(2, sheet.max_row + 1):
        cell_value = sheet.cell(row=row, column=1).value
        if ObjUtil.is_empty(cell_value):
            break
        if cell_value == '删卡':
            for col in range(9, 9 + 2):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SR']['{cell_value}'][{col - 9}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
            for col in range(9 + 3, 9 + 5):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SR']['次限{cell_value}'][{col - 9 - 3}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
        elif cell_value == '百分比' or cell_value == '基础值':
            for col in range(9, 9 + 5):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SR']['{cell_value}'][{col - 9}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
            for col_left in range(9, 9 + 5):
                for col_right in range(col_left + 1, 9 + 5):
                    name1 = ExcelUtil.col_row_to_excel_col(col_left, row)
                    name2 = ExcelUtil.col_row_to_excel_col(col_right, row)
                    value = f"entry['SR']['{cell_value}'][{col_left - 9}:{col_right - 9+1}]"
                    route_dict.update({f"词条!{name1}:{name2}": f"{value}"})
                    route_dict.update({f"'词条'!{name1}:{name2}": f"{value}"})
        else:
            for col in range(9, 9 + 5):
                name = ExcelUtil.col_row_to_excel_col(col, row)
                value = f"entry['SR']['{cell_value}'][{col - 9}]"
                route_dict.update({f"词条!{name}": value})
                route_dict.update({f"'词条'!{name}": value})
    # 保存为JSON文件
    with open(f'../resource/tmp/entry_dict.json', 'w', encoding='utf-8') as f:
        json.dump(route_dict, f, ensure_ascii=False, indent=4)


if __name__ == '__main__':
    func_dict()
    func()
