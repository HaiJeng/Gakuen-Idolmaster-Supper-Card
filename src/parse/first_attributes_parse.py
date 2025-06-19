import json
from tkinter.messagebox import RETRY

import openpyxl

from common.config import attr_name_list
from parse.util.parse_to_first_attr_dict_util import ParseToFirstAttrDictUtil
from service.CardCalculator import CardCalculator
from util.ObjUtil import ObjUtil
from util.StrUtil import StrUtil

# 加载工作簿
wb = openpyxl.load_workbook('../mnt/支援卡-6.9新卡追加.xlsx')

with open("../resource/tmp/cell_dict.json", 'r', encoding='utf-8') as f:
    cell_dict = json.load(f)
# 加载词条数据
with open('../resource/tmp/entry_dict.json', 'r', encoding='utf-8') as f:
    entry_data = json.load(f)

"""读取 route_dict.json 文件并返回字典"""
with open('../resource/tmp/route_dict.json', 'r', encoding='utf-8') as f:
    route_data = json.load(f)


def get_attribute(sheet_name: str):
    sheet = wb[sheet_name]
    first_attribute = {}
    for row in range(2, sheet.max_row):
        if ObjUtil.is_empty(sheet.cell(row=row, column=1).value):
            break
        tmp_name = sheet.cell(row=row, column=1).value
        parts = tmp_name.split("（")
        card_name = parts[0].strip()
        card_value = {}
        calculator_attr_col = 4
        for column in range(4, sheet.max_column):
            if ObjUtil.is_empty(sheet.cell(row=1, column=column).value):
                calculator_attr_col = column
                break
            cell = sheet.cell(row=row, column=column).value
            if StrUtil.is_str(cell):
                cell = cell.replace('ROW()', row.__str__())
            value_name = sheet.cell(row=1, column=column).value
            card_value[value_name] = CardCalculator.convert_formula(cell)
        calculator_attr = {}
        for column in range(calculator_attr_col, sheet.max_column):
            if ObjUtil.is_empty(sheet.cell(row=1, column=column).value):
                continue
            cell = sheet.cell(row=row, column=column).value
            if StrUtil.is_str(cell):
                cell = cell.replace('ROW()', row.__str__())
            value_name = sheet.cell(row=1, column=column).value
            calculator_attr[value_name] = CardCalculator.convert_formula(cell)
        card_value['calculator_attr'] = calculator_attr
        first_attribute[card_name] = card_value
    return first_attribute


def formate_attribute_dict(attr_data: dict, attr_dict: dict):
    json_str = json.dumps(attr_data, ensure_ascii=False)

    # 2. 按键长度降序排序，避免短键误替换长键（如先替换"'持有情况'!B11"再替换"'持有情况'!B1"）
    cell_dict_keys = sorted(cell_dict.keys(), key=len, reverse=True)
    # 2. 按键长度降序排序，避免短键误替换长键（如先替换"'路线'!B11"再替换"'路线'!B1"）
    route_sorted_keys = sorted(route_data.keys(), key=len, reverse=True)
    # 2. 按键长度降序排序，避免短键误替换长键（如先替换"'词条'!B11"再替换"'词条'!B1"）
    entry_sorted_keys = sorted(entry_data.keys(), key=len, reverse=True)
    # # 2. 按键长度降序排序，避免短键误替换长键（如先替换"'词条'!B11"再替换"'词条'!B1"）
    # attr_sorted_keys = sorted(attr_dict.keys(), key=len, reverse=True)
    # 3. 逐个替换Excel引用为中文描述
    for excel_ref in cell_dict_keys:
        json_str = json_str.replace(f"持有情况!{excel_ref.__str__()}", cell_dict[excel_ref.__str__()])
        json_str = json_str.replace(f"'持有情况'!{excel_ref.__str__()}", cell_dict[excel_ref.__str__()])
    # 3. 逐个替换Excel引用为中文描述
    for excel_ref in route_sorted_keys:
        json_str = json_str.replace(excel_ref.__str__(), route_data[excel_ref.__str__()])
    # 3. 逐个替换Excel引用为中文描述
    for excel_ref in entry_sorted_keys:
        json_str = json_str.replace(excel_ref.__str__(), entry_data[excel_ref.__str__()])
    # 3. 逐个替换Excel引用为中文描述
    for excel_ref in attr_dict.keys():
        json_str = json_str.replace(excel_ref.__str__(), attr_dict[excel_ref.__str__()])
    return json.loads(json_str)


def formate_attribute_dict_calculator(attr_dict: dict):
    for key, value in attr_dict.items():
        for attr in value.keys():
            if StrUtil.is_str(value[attr]):
                value[attr] = CardCalculator.convert_formula(value[attr])
    return attr_dict


if __name__ == '__main__':
    for _str in attr_name_list:
        _attr_dict = ParseToFirstAttrDictUtil(wb[_str], _str).parse_to_first_attr_dict_util()
        data = get_attribute(_str)
        data = formate_attribute_dict(data, _attr_dict)
        with open(f'../resource/attr/{_str}.json', 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=4)
    ParseToFirstAttrDictUtil(wb['第一属性比较'], '第一属性比较').parse_to_attr_cell_dict_util()
