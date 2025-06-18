import json

import openpyxl

from model.SupportCard import Rarity
from parse.util.parse_to_cell_dict_util import ParseSheetUtil
from service.CardCalculator import CardCalculator
from util.StrUtil import StrUtil

"""
用于解析excel的持有情况表，并替换公式，数据保存至resource/card/持有情况_数值和公式
"""

# 加载工作簿
wb = openpyxl.load_workbook('../mnt/支援卡-6.9新卡追加.xlsx')
# 加载词条数据
with open('../resource/tmp/entry_dict.json', 'r', encoding='utf-8') as f:
    entry_data = json.load(f)

"""读取 route_dict.json 文件并返回字典"""
with open('../resource/tmp/route_dict.json', 'r', encoding='utf-8') as f:
    route_data = json.load(f)

# 获取指定工作表
sheet = wb['持有情况']


def process_excel_col(row: int, col: int, rarity: str, color: str) -> {}:
    row_max = row + 4
    col_max = col + 7
    tmp_name = sheet.cell(row=row, column=col).value
    # print(row, col, tmp_name)
    parts = tmp_name.split("（")
    cel_data = {
        "name": parts[0].strip(),
        "nickname": parts[1].replace("）", "").strip() if len(parts) > 1 else "",
        "稀有度": rarity,
        "破数": sheet.cell(row=row + 2, column=col + 2).value,
        "attributes": [
            {"第一属性": CardCalculator.convert_formula(sheet.cell(row=row + 2, column=col_max - 2).value)},
            {"第二属性": CardCalculator.convert_formula(sheet.cell(row=row + 2, column=col_max - 1).value)},
            {"第三属性": CardCalculator.convert_formula(sheet.cell(row=row + 2, column=col_max).value)}
        ],
        "tags": [
            val for val in (
                sheet.cell(row=row_max, column=col_max - 2).value,
                sheet.cell(row=row + 3, column=col_max - 2).value,
            ) if val is not None],
        "color": color,
        "bonuses": []
    }
    bonus_candidates = [
        {sheet.cell(row=row + 1, column=col + 3).value: sheet.cell(row=row + 2, column=col + 3).value},
        {sheet.cell(row=row + 1, column=col + 4).value: sheet.cell(row=row + 2, column=col + 4).value},
        {sheet.cell(row=row + 3, column=col + 2).value: sheet.cell(row=row_max, column=col + 2).value},
        {sheet.cell(row=row + 3, column=col + 3).value: sheet.cell(row=row_max, column=col + 3).value},
        {sheet.cell(row=row + 3, column=col + 4).value: sheet.cell(row=row_max, column=col + 4).value}
    ]

    for bonus in bonus_candidates:
        if not any(key is None for key in bonus):
            processed_bonus = {key: CardCalculator.convert_formula(value) for key, value in bonus.items()}
            cel_data["bonuses"].append(processed_bonus)
    # 判断是否有道具
    item = sheet.cell(row=row, column=col_max + 1).value
    if item:
        item_row, item_col = row, col_max + 1
        item = {
            item: sheet.cell(row=item_row + 1, column=item_col).value,
            "道具第一属性加成": sheet.cell(row=item_row + 2, column=item_col).value,
            "道具第二属性加成": sheet.cell(row=item_row + 3, column=item_col).value,
            "道具第三属性加成": sheet.cell(row=item_row + 4, column=item_col).value
        }
        cel_data['item'] = {}
        for key, value in item.items():
            if StrUtil.is_str(value) and value.startswith('='):
                cel_data['item'][key] = CardCalculator.convert_formula(value)
            else:
                cel_data['item'][key] = value
    return cel_data


def process_excel():
    colors = ["红", "蓝", "黄"]
    cols = []
    data = []
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(row=2, column=col).value
        if type(value) is str and StrUtil.contains_japanese(value):
            if col not in cols:  # 避免重复
                cols.append(col)  # 保持插入顺序
    for i in range(0, len(cols)):
        for row in range(1, sheet.max_row + 1):
            col = cols[i]
            value = sheet.cell(row=row, column=col).value
            if type(value) is str and value.__contains__('（'):
                color = colors[i % 3]
                rarity = Rarity.SSR.value
                if i > 3:
                    rarity = Rarity.SR.value
                data.append(process_excel_col(row, col, rarity, color))
    s = replace_refs_with_dict(data)
    # 保存结果
    with open('../resource/card/持有情况_数值和公式.json', 'w', encoding='utf-8') as f:
        json.dump(s, f, ensure_ascii=False, indent=4)


def replace_refs_with_dict(data: list) -> list:
    # 1. 将data转为JSON字符串
    json_str = json.dumps(data, ensure_ascii=False)
    print("route_data.keys()", route_data.keys())
    # 2. 按键长度降序排序，避免短键误替换长键（如先替换"'路线'!B11"再替换"'路线'!B1"）
    route_sorted_keys = sorted(route_data.keys(), key=len, reverse=True)
    print("entry_data.keys()", entry_data.keys())
    # 2. 按键长度降序排序，避免短键误替换长键（如先替换"'词条'!B11"再替换"'词条'!B1"）
    entry_sorted_keys = sorted(entry_data.keys(), key=len, reverse=True)

    # 3. 逐个替换Excel引用为中文描述
    for excel_ref in route_sorted_keys:
        json_str = json_str.replace(excel_ref.__str__(), route_data[excel_ref.__str__()])
    # 4. 逐个替换Excel引用为中文描述
    for excel_ref in entry_sorted_keys:
        json_str = json_str.replace(excel_ref.__str__(), entry_data[excel_ref.__str__()])
    # 5. 替换在持有情况页json和单元格映射
    all_cell_dict = ParseSheetUtil(sheet).process_excel()
    all_cell_sorted_keys = sorted(all_cell_dict.keys(), key=len, reverse=True)
    for excel_ref in all_cell_sorted_keys:
        json_str = json_str.replace(excel_ref.__str__(), all_cell_dict[excel_ref.__str__()])

    # 4. 将字符串转回Python对象
    return json.loads(json_str)


if __name__ == '__main__':
    process_excel()
