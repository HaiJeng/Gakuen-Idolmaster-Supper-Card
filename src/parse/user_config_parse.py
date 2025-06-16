import json
from typing import Dict

import openpyxl

from common.config import catch_card
from util.ObjUtil import ObjUtil

wb = openpyxl.load_workbook('../mnt/支援卡-6.9新卡追加.xlsx')
# 获取路线工作表
sheet = wb['路线']


def parse_config():
    user_config: Dict[str, float] = dict()
    row_max = 0
    for row in range(1, sheet.max_row):
        name = sheet.cell(row=row, column=1).value
        if ObjUtil.is_empty(name):
            row_max = row - 1
            break
    for row in range(1, row_max):
        name = sheet.cell(row=row, column=1).value
        value = sheet.cell(row=row, column=2).value
        user_config.update({name: value})
    user_config[sheet.cell(row=row_max, column=1).value] = sheet.cell(row=row_max, column=3).value
    user_config[sheet.cell(row=1, column=3).value] = {
        "红": sheet.cell(row=2, column=3).value,
        "蓝": sheet.cell(row=3, column=3).value,
        "黄": sheet.cell(row=4, column=3).value
    }
    for key in catch_card:
        user_config[key] = 10
    with open(f'../resource/user_config.json', 'w', encoding='utf-8') as f:
        json.dump(user_config, f, ensure_ascii=False, indent=4)


if __name__ == '__main__':
    parse_config()
