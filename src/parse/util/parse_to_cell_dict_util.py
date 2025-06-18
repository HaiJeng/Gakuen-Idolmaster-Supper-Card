import json

import openpyxl

from util.ExcelUtil import ExcelUtil
from util.ObjUtil import ObjUtil
from util.StrUtil import StrUtil


class ParseSheetUtil:
    def __init__(self, sheet):
        self.sheet = sheet

    def process_excel(self) -> {}:
        cols = []
        for col in range(1, self.sheet.max_column + 1):
            value = self.sheet.cell(row=2, column=col).value
            if type(value) is str and StrUtil.contains_japanese(value):
                if col not in cols:  # 避免重复
                    cols.append(col)  # 保持插入顺序
        cell_dict = {}
        for i in range(0, len(cols)):
            for row in range(1, self.sheet.max_row + 1):
                col = cols[i]
                value = self.sheet.cell(row=row, column=col).value
                if type(value) is str and StrUtil.contains_japanese(value):
                    cell_dict.update(self.process_excel_col(row, col))
        return cell_dict

    def process_excel_col(self, row: int, col: int) -> {}:
        cel_data = {}
        col_max = col + 7
        cel_data.update({ExcelUtil.row_col_to_excel_col(row + 2, col + 2): "card['破数']"})
        cel_data.update({ExcelUtil.row_col_to_excel_col(row + 2, col_max - 2): "card['第一属性']"})
        cel_data.update({ExcelUtil.row_col_to_excel_col(row + 2, col_max - 1): "card['第二属性']"})
        cel_data.update({ExcelUtil.row_col_to_excel_col(row + 2, col_max): "card['第三属性']"})
        bonus_candidates = [
            {ExcelUtil.row_col_to_excel_col(row + 2, col + 3): self.sheet.cell(row=row + 1, column=col + 3).value},
            {ExcelUtil.row_col_to_excel_col(row + 2, col + 4): self.sheet.cell(row=row + 1, column=col + 4).value},
            {ExcelUtil.row_col_to_excel_col(row + 4, col + 2): self.sheet.cell(row=row + 3, column=col + 2).value},
            {ExcelUtil.row_col_to_excel_col(row + 4, col + 3): self.sheet.cell(row=row + 3, column=col + 3).value},
            {ExcelUtil.row_col_to_excel_col(row + 4, col + 4): self.sheet.cell(row=row + 3, column=col + 4).value}
        ]
        for bonus in bonus_candidates:
            if not any(value is None for value in bonus.values()):
                processed_bonus = {key: f"card['{value}']" for key, value in bonus.items()}
                cel_data.update(processed_bonus)
        # 判断是否有道具
        item = self.sheet.cell(row=row, column=col_max + 1).value
        if item:
            item_row, item_col = row, col_max + 1
            cel_data.update({ExcelUtil.row_col_to_excel_col(item_row + 1, item_col): "card['道具属性']"})
            cel_data.update({ExcelUtil.row_col_to_excel_col(item_row + 2, item_col): "card['道具第一属性加成']"})
            cel_data.update({ExcelUtil.row_col_to_excel_col(item_row + 3, item_col): "card['道具第二属性加成']"})
            cel_data.update({ExcelUtil.row_col_to_excel_col(item_row + 4, item_col): "card['道具第三属性加成']"})
            user_value = self.sheet.cell(row=item_row, column=item_col + 1).value
            if ObjUtil.not_empty(user_value):
                cel_data.update({
                    ExcelUtil.row_col_to_excel_col(item_row + 1, item_col + 1): f"user['{user_value}']",
                })
        return cel_data


# if __name__ == '__main__':
# wb = openpyxl.load_workbook('../mnt/支援卡-6.9新卡追加.xlsx')
# sheet = wb['持有情况']
# cell_dict = ParseSheetUtil(sheet).process_excel()
#
# with open(f'../resource/tmp/cell_dict.json', 'w', encoding='utf-8') as _f:
#     json.dump(cell_dict, _f, ensure_ascii=False, indent=4)
