import re
from openpyxl import load_workbook
import json

# 加载Excel文件
wb = load_workbook('../mnt/支援卡-6.9新卡追加.xlsx')

# 获取路线工作表
sheet = wb['路线']


def func(route: str):
    # 查找route单元格的位置
    route_choice_row = None
    route_choice_col = None

    for row in range(1, sheet.max_row + 1):
        for col in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=row, column=col).value
            if cell_value == route:
                route_choice_row = row
                route_choice_col = col
                break
        if route_choice_row:
            break
    # 获取Key
    routes_name_len = []
    if route_choice_row and route_choice_col:
        for row in range(route_choice_row + 1, sheet.max_row + 1):
            route_name = sheet.cell(row=row, column=1).value
            if not route_name:
                break  # 遇到空行时停止
            routes_name_len.append(route_name)
    print(routes_name_len)

    # 解析IF公式的正则表达式
    if_pattern = re.compile(r'=IF\((.*?),(.*?),(.*?)\)')

    # 初始化两个结果字典
    result_true = {}  # 条件为真时的值
    result_false = {}  # 条件为假时的值
    # 提取路线选择数据
    routes_data = {}
    if route_choice_row and route_choice_col:
        # 指定路线下方的单元格
        attribute_len = []
        for row in range(route_choice_row + 1, sheet.max_row + 1):
            attribute = sheet.cell(row=row, column=route_choice_col).value
            if attribute is None or attribute == '':
                break  # 遇到空行时停止
            attribute_len.append(attribute)
            print(attribute)
        print(attribute_len, attribute_len.__len__())
        has_if = False
        for value in attribute_len:
            if isinstance(value, str) and value.startswith('=IF('):
                has_if = True
                break
        if has_if:
            for i in range(0, attribute_len.__len__()):
                value = attribute_len[i]
                if isinstance(value, str) and value.startswith('=IF('):
                    match = if_pattern.match(value)
                    if match:
                        condition = match.group(1)
                        true_value = match.group(2)
                        false_value = match.group(3)

                        # 尝试转换为数字，否则保留字符串
                        try:
                            true_value = float(true_value) if '.' in true_value else int(true_value)
                        except ValueError:
                            true_value = true_value.strip('"')

                        try:
                            false_value = float(false_value) if '.' in false_value else int(false_value)
                        except ValueError:
                            false_value = false_value.strip('"')

                        # 存储到对应的结果字典
                        result_true[routes_name_len[i]] = true_value
                        result_false[routes_name_len[i]] = false_value
                else:
                    # 非IF公式的值直接复制到两个结果字典
                    result_true[routes_name_len[i]] = value
                    result_false[routes_name_len[i]] = value
            # 保存为JSON文件
            with open(f'../resource/route/{route}_单极.json', 'w', encoding='utf-8') as f:
                json.dump(result_true, f, ensure_ascii=False, indent=4)

            with open(f'../resource/route/{route}_均衡.json', 'w', encoding='utf-8') as f:
                json.dump(result_false, f, ensure_ascii=False, indent=4)
        else:
            for i in range(0, attribute_len.__len__()):
                value = attribute_len[i]
                routes_data[routes_name_len[i]] = value
            # 保存为JSON文件
            with open(f'../resource/route/{route}.json', 'w', encoding='utf-8') as f:
                json.dump(routes_data, f, ensure_ascii=False, indent=4)


def func_dict():
    # 查找route单元格的位置
    route_choice_row = None

    for row in range(1, sheet.max_row + 1):
        cell_value = sheet.cell(row=row, column=2).value
        if cell_value == '路线选择':
            route_choice_row = row
            break
    route_dict = {}
    for row in range(route_choice_row + 1, sheet.max_row + 1):
        value = sheet.cell(row=row, column=1).value
        print(value)
        if not value:
            break
        route_dict.update({'\'路线\'!B'+str(row): value})
    print(route_dict)
    # 保存为JSON文件
    with open(f'../resource/route/route_dict.json', 'w', encoding='utf-8') as f:
        json.dump(route_dict, f, ensure_ascii=False, indent=4)

if __name__ == '__main__':
    func_dict()
    func("路线选择")
    func("NIA通常")
    func("NIA强化")
    func("NIA购物")
    func("自定义")
    func("mas支给")
    func("mas外出")
    func("mas购物")
    func("mas买卡")
    func("mas删卡")
    func("mas外出强化")
